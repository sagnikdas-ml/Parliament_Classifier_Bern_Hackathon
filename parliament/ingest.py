import csv
import json
import re
from html.parser import HTMLParser
from itertools import islice

import pyarrow as pa
from openpyxl import load_workbook

from parliament import db
from parliament.config import API, DATA, FILES, LANGUAGES, digest, json_text
from parliament.download import api_rows, bootstrap_files, cache_json, download, open_url, read_records


class TextParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"p", "br", "div", "li", "tr"}:
            self.parts.append(" ")

    def handle_endtag(self, tag):
        if tag in {"p", "div", "li", "tr"}:
            self.parts.append(" ")

    def handle_data(self, data):
        self.parts.append(data)


def clean(text):
    parser = TextParser()
    parser.feed(str(text or ""))
    return " ".join("".join(parser.parts).split())


def translations(record, field):
    nested = record.get(field)
    if isinstance(nested, dict):
        return {language: clean(value) for language, value in nested.items() if language in LANGUAGES and value}
    return {language: clean(record[f"{field}_{language}"]) for language in LANGUAGES if record.get(f"{field}_{language}")}


def unit(affair_id, identifier, language, kind, text, url):
    text = clean(text)
    return {"id": str(identifier), "affair_id": int(affair_id), "language": language,
            "kind": kind, "text": text, "source_url": url, "content_hash": digest(text)}


def upsert(connection, table, records):
    if not records:
        return
    incoming = pa.Table.from_pylist(records)
    columns = ",".join(records[0])
    connection.register("incoming_records", incoming)
    try:
        connection.execute(f"INSERT OR REPLACE INTO {table} ({columns}) SELECT {columns} FROM incoming_records")
    finally:
        connection.unregister("incoming_records")


def affair_record(record):
    identifier = int(record.get("id") or record["openparldata_id"])
    titles = translations(record, "title")
    url = f"{API}/affairs/{identifier}"
    affair = {"id": identifier, "body_key": str(record.get("body_key") or ""),
              "begin_date": record.get("begin_date"), "titles": json_text(titles),
              "source_url": url, "source_hash": digest(record), "raw": json_text(record)}
    units = [unit(identifier, f"title:{identifier}:{language}", language, "title", title, url)
             for language, title in titles.items()]
    return affair, units


def import_affairs(path):
    source = iter(read_records(path))
    count = 0
    while batch := list(islice(source, 2000)):
        affairs, units = [], []
        for record in batch:
            affair, title_units = affair_record(record)
            affairs.append(affair)
            units.extend(title_units)
        with db.connect() as connection:
            connection.execute("BEGIN")
            try:
                connection.execute("DELETE FROM units WHERE kind='title' AND affair_id IN (SELECT unnest(?::BIGINT[]))", [[affair["id"] for affair in affairs]])
                upsert(connection, "affairs", affairs)
                upsert(connection, "units", units)
                connection.execute("COMMIT")
            except Exception:
                connection.execute("ROLLBACK")
                raise
        count += len(batch)
    return count


def import_topics(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        records = workbook["Topics"].iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(records)]
        topics = []
        for values in records:
            record = dict(zip(headers, values))
            if not record.get("code"):
                continue
            name = " / ".join(str(record.get(field) or "") for field in ("topic_de", "topic_fr"))
            definition = " | ".join(str(record.get(field) or "") for field in ("chapter_de", "chapter_fr", "topic_de", "topic_fr"))
            topics.append({"id": f"aps:{record['code']}", "scheme": "aps", "name": name,
                           "definition": definition, "version": 1, "raw": json_text(record)})
        if not topics:
            raise ValueError("APS workbook contains no recognised topic rows")
        with db.connect() as connection:
            upsert(connection, "topics", topics)
        return len(topics)
    finally:
        workbook.close()


def import_labels(path):
    name = path.name
    if name.endswith(".csv"):
        with path.open(encoding="utf-8-sig", newline="") as source:
            records = list(csv.DictReader(source))
        scheme = "politmonitor:digitalisation" if name.startswith("politmonitor") else "dvs:digitalisation"
        values = [{"affair_id": int(record["openparldata_id"]), "scheme": scheme, "labels": json_text(record)} for record in records if record.get("openparldata_id")]
    else:
        values = []
        for record in read_records(path):
            scheme = "curia_vista" if name.startswith("che_") else f"source:{record.get('body_key', 'unknown')}"
            values.append({"affair_id": int(record["openparldata_id"]), "scheme": scheme, "labels": json_text(record)})
    with db.connect() as connection:
        upsert(connection, "source_labels", values)


def bootstrap(refresh=False):
    result = {}
    for path in bootstrap_files(refresh):
        if path.name == "affairs.ndjson.gz":
            result["affairs"] = import_affairs(path)
        elif path.name == "bodies.ndjson.gz":
            with db.connect() as connection:
                upsert(connection, "bodies", [{"id": str(record.get("id") or record.get("key")), "raw": json_text(record)} for record in read_records(path)])
        elif path.suffix == ".xlsx":
            result["topics"] = import_topics(path)
        elif path.suffix in (".csv", ".jsonl"):
            import_labels(path)
    return result


def ensure_affair(affair_id, refresh=False):
    existing = db.query("SELECT * FROM affairs WHERE id=?", [affair_id])
    if existing and not refresh:
        return existing[0]
    affair, units = affair_record(cache_json(f"{API}/affairs/{affair_id}", refresh))
    with db.connect() as connection:
        upsert(connection, "affairs", [affair])
        connection.execute("DELETE FROM units WHERE affair_id=? AND kind='title'", [affair_id])
        upsert(connection, "units", units)
    return affair


def text_units(record, docs=False):
    affair_id = record.get("affair_id") or record.get("openparldata_id")
    if not affair_id:
        return []
    identifier = record.get("id") or record.get("text_id")
    kind = "document" if docs else "text"
    url = record.get("url") if docs else f"{API}/affairs/{affair_id}/texts"
    values = translations(record, "text")
    if docs and isinstance(record.get("text"), str):
        language = str(record.get("language") or "und").lower()
        values = {language: record["text"]}
    return [unit(affair_id, f"{kind}:{affair_id}:{identifier}:{language}", language, kind, value, url or f"{API}/affairs/{affair_id}/docs") for language, value in values.items() if clean(value)]


def fetch_content(affair_id, docs=False, refresh=False):
    ensure_affair(affair_id, refresh)
    kind = "docs" if docs else "texts"
    if not refresh and db.query("SELECT 1 FROM fetches WHERE affair_id=? AND kind=?", [affair_id, kind]):
        return
    records = list(api_rows(f"{API}/affairs/{affair_id}/{kind}"))
    units = [item for record in records for item in text_units(record, docs)]
    with db.connect() as connection:
        connection.execute("BEGIN")
        try:
            connection.execute("DELETE FROM units WHERE affair_id=? AND kind=?", [affair_id, "document" if docs else "text"])
            upsert(connection, "units", units)
            connection.execute("INSERT OR REPLACE INTO fetches (affair_id, kind) VALUES (?, ?)", [affair_id, kind])
            connection.execute("COMMIT")
        except Exception:
            connection.execute("ROLLBACK")
            raise


def import_texts(path, docs=False):
    source = iter(read_records(path))
    count = 0
    while batch := list(islice(source, 1000)):
        units = [item for record in batch for item in text_units(record, docs)]
        with db.connect() as connection:
            upsert(connection, "units", units)
        count += len(units)
    return count


def bulk_texts(docs=False, refresh=False):
    folder = "docs" if docs else "texts"
    url = f"{FILES}/exports/{folder}/"
    with open_url(url) as response:
        listing = response.read().decode()
    files = re.findall(r'href="([A-Za-z0-9_.-]+\.ndjson\.gz)"', listing)
    if not files:
        raise ValueError("No export partitions found")
    result = {}
    for name in files:
        print(f"Preparing {name}", flush=True)
        path = download(url + name, DATA / "raw" / "exports" / folder / name, refresh)
        result[name] = import_texts(path, docs)
    return result


def sample_affairs(limit=200):
    candidates = db.query("""
        WITH candidates AS (
          SELECT a.id, a.body_key, u.language,
            CASE WHEN a.body_key='CHE' THEN 'federal'
                 WHEN regexp_matches(a.body_key, '^[A-Z]{2}$') THEN 'cantonal'
                 ELSE 'municipal' END AS level,
            row_number() OVER (PARTITION BY a.id ORDER BY hash(u.language || CAST(a.id AS VARCHAR))) AS choice
          FROM affairs a JOIN units u ON a.id=u.affair_id
          WHERE u.kind='title' AND u.language IN ('de','fr','it')
            AND a.body_key <> 'LIE' AND a.begin_date >= '2018'
        ), ranked AS (
          SELECT *, row_number() OVER (PARTITION BY language,level ORDER BY hash(id)) AS position
          FROM candidates WHERE choice=1
        ) SELECT id AS affair_id, language, body_key, level FROM ranked
        ORDER BY position,language,level LIMIT ?
    """, [limit])
    return candidates
